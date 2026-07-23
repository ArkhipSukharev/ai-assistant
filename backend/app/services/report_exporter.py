from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FONT_PATHS = (
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
)


def _display(value: Any, value_format: str | None = None) -> str:
    if value is None:
        return "—"
    if value_format == "currency":
        return f"{float(value):,.0f} ₽".replace(",", " ")
    if value_format == "percent":
        return f"{float(value):g}%"
    if value_format == "boolean":
        return "Да" if value else "Нет"
    return str(value)


def render_xlsx(presentation: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"
    sheet.sheet_view.showGridLines = False

    sheet["A1"] = presentation.get("title", "Отчёт")
    sheet["A1"].font = Font(size=18, bold=True, color="173F35")
    row = 3
    for metric in presentation.get("metrics", []):
        sheet.cell(row=row, column=1, value=metric.get("label"))
        sheet.cell(
            row=row,
            column=2,
            value=_display(metric.get("value"), metric.get("format")),
        )
        sheet.cell(row=row, column=1).font = Font(bold=True)
        row += 1

    plan = presentation.get("plan")
    if plan:
        row += 1
        sheet.cell(row=row, column=1, value="План / факт").font = Font(bold=True)
        row += 1
        for label, value in (
            ("План", plan.get("target")),
            ("Факт", plan.get("actual")),
            ("Выполнение", f"{plan.get('completion_percent', 0)}%"),
        ):
            sheet.cell(row=row, column=1, value=label)
            sheet.cell(row=row, column=2, value=value)
            row += 1

    table = presentation.get("table") or {}
    columns = table.get("columns", [])
    if columns:
        row += 1
        header_row = row
        for column_index, column in enumerate(columns, 1):
            cell = sheet.cell(
                row=header_row,
                column=column_index,
                value=column.get("label", column.get("key")),
            )
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="10A37F")
            cell.alignment = Alignment(vertical="center")
        for item in table.get("rows", []):
            row += 1
            for column_index, column in enumerate(columns, 1):
                sheet.cell(
                    row=row,
                    column=column_index,
                    value=_display(item.get(column["key"]), column.get("format")),
                )

    provenance = presentation.get("provenance") or {}
    row += 2
    sheet.cell(
        row=row,
        column=1,
        value=(
            f"Источник: {provenance.get('source', 'amoCRM')} · "
            f"{provenance.get('fetched_at', '')}"
        ),
    )
    sheet.cell(row=row, column=1).font = Font(size=9, color="777777")

    for index in range(1, max(len(columns), 2) + 1):
        values = [
            len(str(sheet.cell(row=item, column=index).value or ""))
            for item in range(1, sheet.max_row + 1)
        ]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(values, default=10) + 3, 45
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_font() -> str:
    for path in FONT_PATHS:
        if path.exists():
            name = "F5ReportFont"
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
            return name
    return "Helvetica"


def render_pdf(presentation: dict[str, Any]) -> bytes:
    output = io.BytesIO()
    font = _pdf_font()
    styles = getSampleStyleSheet()
    for style_name in ("Title", "Heading2", "BodyText"):
        styles[style_name].fontName = font

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=presentation.get("title", "Отчёт"),
    )
    story = [Paragraph(presentation.get("title", "Отчёт"), styles["Title"]), Spacer(1, 5 * mm)]

    metrics = presentation.get("metrics", [])
    if metrics:
        metric_data = [
            [
                Paragraph(str(metric.get("label", "")), styles["BodyText"]),
                Paragraph(
                    _display(metric.get("value"), metric.get("format")),
                    styles["BodyText"],
                ),
            ]
            for metric in metrics
        ]
        metric_table = Table(metric_data, colWidths=[55 * mm, 45 * mm])
        metric_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F0F3F2")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9DEDC")),
                    ("PADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.extend([metric_table, Spacer(1, 6 * mm)])

    plan = presentation.get("plan")
    if plan:
        story.extend(
            [
                Paragraph(
                    "План / факт: "
                    f"{_display(plan.get('actual'))} из {_display(plan.get('target'))} "
                    f"({plan.get('completion_percent', 0)}%)",
                    styles["Heading2"],
                ),
                Spacer(1, 4 * mm),
            ]
        )

    table = presentation.get("table") or {}
    columns = table.get("columns", [])
    rows = table.get("rows", [])
    if columns:
        data = [
            [
                Paragraph(str(column.get("label", column["key"])), styles["BodyText"])
                for column in columns
            ]
        ]
        for item in rows:
            data.append(
                [
                    Paragraph(
                        _display(item.get(column["key"]), column.get("format")),
                        styles["BodyText"],
                    )
                    for column in columns
                ]
            )
        width = 260 * mm / max(len(columns), 1)
        report_table = Table(data, repeatRows=1, colWidths=[width] * len(columns))
        report_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), font),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10A37F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9DEDC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(report_table)

    provenance = presentation.get("provenance") or {}
    story.extend(
        [
            Spacer(1, 5 * mm),
            Paragraph(
                f"Источник: {provenance.get('source', 'amoCRM')} · "
                f"{provenance.get('fetched_at', '')}",
                styles["BodyText"],
            ),
        ]
    )
    document.build(story)
    return output.getvalue()
